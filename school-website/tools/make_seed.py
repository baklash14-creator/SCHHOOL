"""Rebuild data/seed.json (demo data) from the spreadsheet. Run from the site folder:
   pip install openpyxl && python tools/make_seed.py"""
import json, datetime as dt, sys
from openpyxl import load_workbook
SRC = sys.argv[1] if len(sys.argv) > 1 else 'spreadsheet/School_Management_v11.xlsx'
v = load_workbook(SRC, data_only=True); f = load_workbook(SRC)
DATA = ['Schools','Classes','Teachers','Students','Attendance','Grades','Fees','Timetable','Staff Attendance','Payroll','Email Log','Report Files']
fx = lambda x: x.strftime('%Y-%m-%d') if isinstance(x, (dt.datetime, dt.date)) else ('' if x is None else x)
sheets = {}
for n in DATA:
    ws, wf = v[n], f[n]; w = ws.max_column
    head = [ws.cell(4, c).value for c in range(1, w + 1)]
    edit = [not str(wf.cell(6, c).value or '').startswith('=') for c in range(1, w + 1)]
    rows = []
    for r in range(5, ws.max_row + 1):
        vals = [fx(ws.cell(r, c).value) for c in range(1, w + 1)]
        if any(x != '' for x in vals): rows.append({'_r': r, **dict(zip(head, vals))})
    sheets[n] = {'headers': head, 'edit': edit, 'rows': rows}
raw = {n: [[fx(c) for c in row] for row in v[n].iter_rows(values_only=True)] for n in ('Guide', 'Settings')}
S = v['Settings']
col = lambda a, b, c: [x for x in (S.cell(r, c).value for r in range(a, b + 1)) if x not in (None, '')]
meta = dict(year=S['B5'].value, currency=S['B6'].value, max=S['B7'].value, cw=S['B8'].value,
            terms=col(11, 13, 1), classes=col(11, 30, 2), subjects=col(11, 25, 3),
            bands=[[S.cell(r, 5).value, S.cell(r, 6).value, S.cell(r, 7).value] for r in range(11, 16)])
json.dump({'meta': meta, 'sheets': sheets, 'raw': raw}, open('data/seed.json', 'w'), default=str)
print(meta)
